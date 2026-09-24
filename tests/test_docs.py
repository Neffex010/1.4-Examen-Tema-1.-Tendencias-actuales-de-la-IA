import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "api"))

pytest = __import__("pytest")

pytest.importorskip("PyPDF2")
pytest.importorskip("docx")
pytest.importorskip("fpdf")
pytest.importorskip("openai")

from docs import DocumentExtractor, MAX_DOC_CHARS  # noqa: E402


def test_extract_from_txt(tmp_path):
    f = tmp_path / "ejemplo.txt"
    f.write_text("  Hola mundo  ", encoding="utf-8")
    assert DocumentExtractor.extract_from_txt(str(f)) == "Hola mundo"


def test_extract_empty_txt(tmp_path):
    f = tmp_path / "vacio.txt"
    f.write_text("", encoding="utf-8")
    assert DocumentExtractor.extract_from_txt(str(f)) == ""


def test_pdf_roundtrip(tmp_path):
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(0, 10, "Hola PDF")
    pdf_path = tmp_path / "sample.pdf"
    pdf.output(str(pdf_path))

    text = DocumentExtractor.extract_from_pdf(str(pdf_path))
    assert "Hola PDF" in text


def test_docx_roundtrip(tmp_path):
    from docx import Document

    d = Document()
    d.add_paragraph("Primer parrafo")
    d.add_paragraph("Segundo parrafo")
    d.add_paragraph("   ")  # debe filtrarse
    docx_path = tmp_path / "sample.docx"
    d.save(str(docx_path))

    text = DocumentExtractor.extract_from_docx(str(docx_path))
    assert text == "Primer parrafo\nSegundo parrafo"


def test_doc_char_limit_sane():
    assert 15000 <= MAX_DOC_CHARS <= 20000