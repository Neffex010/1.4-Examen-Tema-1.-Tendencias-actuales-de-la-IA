import base64
import binascii
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from openai import OpenAI
import PyPDF2
import docx
from docx.shared import Cm
from fpdf import FPDF

from _helpers import (
    BaseTranslatorHandler,
    clean_xml,
    sanitize_latin1,
    valid_extension,
    valid_language,
    ALLOWED_DOC_EXTENSIONS,
)

MAX_DOC_MB = 5
MAX_DOC_BYTES = MAX_DOC_MB * 1024 * 1024
MAX_DOC_CHARS = 15000


class DocumentExtractor:
    @staticmethod
    def extract_from_txt(path):
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read().strip()

    @staticmethod
    def extract_from_pdf(path):
        text = ""
        with open(path, "rb") as f:
            for page in PyPDF2.PdfReader(f).pages:
                text += (page.extract_text() or "") + "\n"
        return text.strip()

    @staticmethod
    def extract_from_docx(path):
        paragraphs = [p.text for p in docx.Document(path).paragraphs if p.text.strip()]
        return "\n".join(paragraphs)


class DocumentExporter:
    @staticmethod
    def to_docx(translated_text, filename, output_path):
        doc_out = docx.Document()
        try:
            doc_out.core_properties.title = f"Traduccion - {filename}"
            doc_out.core_properties.author = "Traductor Inteligente IA"
        except Exception:
            pass

        for section in doc_out.sections:
            section.top_margin = Cm(2.5)
            section.bottom_margin = Cm(2.5)
            section.left_margin = Cm(2.5)
            section.right_margin = Cm(2.5)

        doc_out.add_heading("Traduccion", level=0)
        for line in translated_text.split("\n"):
            s = clean_xml(line).strip()
            if not s:
                continue
            if s.startswith("### "):
                doc_out.add_heading(s[4:], level=3)
            elif s.startswith("## "):
                doc_out.add_heading(s[3:], level=2)
            elif s.startswith("# "):
                doc_out.add_heading(s[2:], level=1)
            elif s.startswith("- ") or s.startswith("* "):
                doc_out.add_paragraph(s[2:], style="List Bullet")
            else:
                doc_out.add_paragraph(s)
        doc_out.save(output_path)

    @staticmethod
    def to_pdf(translated_text, output_path):
        class _PDF(FPDF):
            def header(self):
                self.set_font("Helvetica", "B", 9)
                self.set_text_color(120, 120, 120)
                self.cell(0, 8, "Traduccion - Traductor Inteligente IA", 0, 1, "R")
                self.ln(2)
                self.set_x(self.l_margin)
                self.set_text_color(30, 30, 30)
                self.set_font("Helvetica", size=11)

            def footer(self):
                self.set_y(-15)
                self.set_font("Helvetica", "I", 8)
                self.set_text_color(150, 150, 150)
                self.cell(0, 10, f"Pagina {self.page_no()}/{{nb}}", 0, 0, "C")

        pdf = _PDF()
        pdf.alias_nb_pages()
        pdf.set_auto_page_break(auto=True, margin=20)
        pdf.set_margins(20, 20, 20)
        pdf.add_page()

        pdf.set_font("Helvetica", "B", 18)
        pdf.set_text_color(20, 60, 120)
        pdf.set_x(pdf.l_margin)
        pdf.cell(0, 12, sanitize_latin1("Traduccion"), 0, 1)
        pdf.ln(4)
        pdf.set_font("Helvetica", size=11)
        pdf.set_text_color(30, 30, 30)

        epw = pdf.epw
        for line in translated_text.split("\n"):
            s = line.strip()
            pdf.set_x(pdf.l_margin)
            if not s:
                pdf.ln(3)
                continue
            if s.startswith("### "):
                pdf.set_font("Helvetica", "B", 13)
                pdf.ln(2)
                pdf.multi_cell(epw, 7, text=sanitize_latin1(s[4:]))
                pdf.set_font("Helvetica", size=11)
                pdf.ln(1)
            elif s.startswith("## "):
                pdf.set_font("Helvetica", "B", 15)
                pdf.ln(3)
                pdf.multi_cell(epw, 8, text=sanitize_latin1(s[3:]))
                pdf.set_font("Helvetica", size=11)
                pdf.ln(2)
            elif s.startswith("# "):
                pdf.set_font("Helvetica", "B", 16)
                pdf.ln(4)
                pdf.multi_cell(epw, 9, text=sanitize_latin1(s[2:]))
                pdf.set_font("Helvetica", size=11)
                pdf.ln(3)
            elif s.startswith("- ") or s.startswith("* "):
                pdf.multi_cell(epw, 6, text=sanitize_latin1("  - " + s[2:]))
            else:
                pdf.multi_cell(epw, 6, text=sanitize_latin1(s))
        pdf.output(output_path)


class handler(BaseTranslatorHandler):
    def do_POST(self):
        payload, error, status = self._read_json()
        if error:
            self._send_json(status, {"error": error})
            return

        target_lang = payload.get("target_language")
        filename = payload.get("filename") or "documento"

        if not valid_language(target_lang):
            self._send_json(400, {"error": "Idioma destino no valido. Usa 'es' o 'en'."})
            return

        extension = filename.split(".")[-1].lower() if "." in filename else ""
        if not valid_extension(extension, ALLOWED_DOC_EXTENSIONS):
            self._send_json(400, {"error": "Formato no permitido. Usa TXT, PDF o DOCX."})
            return

        try:
            file_bytes = base64.b64decode(payload.get("file") or "", validate=False)
        except (binascii.Error, ValueError):
            self._send_json(400, {"error": "El documento enviado no es Base64 valido."})
            return

        if not file_bytes:
            self._send_json(400, {"error": "El documento esta vacio."})
            return
        if len(file_bytes) > MAX_DOC_BYTES:
            self._send_json(413, {"error": f"El documento excede el limite de {MAX_DOC_MB}MB."})
            return

        with tempfile.NamedTemporaryFile(delete=False, suffix=f".{extension}") as tmp_file:
            tmp_file.write(file_bytes)
            tmp_path = tmp_file.name

        docx_path = f"{tmp_path}_out.docx"
        pdf_path = f"{tmp_path}_out.pdf"
        try:
            extractor = DocumentExtractor()
            if extension == "txt":
                text = extractor.extract_from_txt(tmp_path)
            elif extension == "pdf":
                text = extractor.extract_from_pdf(tmp_path)
            else:
                text = extractor.extract_from_docx(tmp_path)

            if len(text) < 5:
                self._send_json(400, {"error": "Documento sin texto procesable."})
                return
            if len(text) > MAX_DOC_CHARS:
                self._send_json(413, {
                    "error": f"El documento es demasiado largo ({len(text)} caracteres). "
                             f"El limite es de {MAX_DOC_CHARS} caracteres."
                })
                return

            client = OpenAI(api_key=os.getenv("OPENAI_API_KEY"))
            trans_res = client.chat.completions.create(
                model="gpt-4o-mini",
                messages=[
                    {
                        "role": "user",
                        "content": f"Traduce al {target_lang.lower()} manteniendo el formato exacto:\n\n{text}",
                    }
                ],
            )
            translated_text = trans_res.choices[0].message.content

            DocumentExporter.to_docx(translated_text, filename, docx_path)
            with open(docx_path, "rb") as f:
                docx_b64 = base64.b64encode(f.read()).decode("utf-8")

            DocumentExporter.to_pdf(translated_text, pdf_path)
            with open(pdf_path, "rb") as f:
                pdf_b64 = base64.b64encode(f.read()).decode("utf-8")

            self._send_json(200, {
                "original_text": text,
                "translated_text": translated_text,
                "docx_b64": docx_b64,
                "pdf_b64": pdf_b64,
            })
        except ValueError as exc:
            self._send_json(400, {"error": str(exc)})
        except Exception as exc:
            print(f"[docs] error interno: {exc}", flush=True)
            self._send_json(500, {"error": "Error al procesar el documento. Intentalo de nuevo."})
        finally:
            for path in (tmp_path, docx_path, pdf_path):
                if os.path.exists(path):
                    os.remove(path)